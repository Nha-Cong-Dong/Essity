"""Small DOCX and PPTX files built in the test, so no binary fixture is committed (T7.3)."""

# pyright: reportMissingTypeStubs=false, reportUnknownMemberType=false
# pyright: reportUnknownVariableType=false, reportUnknownArgumentType=false
# pyright: reportPrivateUsage=false, reportAttributeAccessIssue=false
# pyright: reportOptionalMemberAccess=false

import io
import zipfile
from collections.abc import Sequence

from docx import Document
from lxml import etree
from pptx import Presentation

P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def make_docx(items: Sequence[tuple[int | None, str]], *, outline_lvl: bool = False) -> bytes:
    """items: (heading level starting at 1, text) or (None, body text). With `outline_lvl`
    a heading is a plain paragraph carrying w:outlineLvl instead of a Heading style."""
    document = Document()
    for level, text in items:
        if level is None:
            document.add_paragraph(text)
        elif outline_lvl:
            paragraph = document.add_paragraph(text)
            p_pr = paragraph._p.get_or_add_pPr()
            node = etree.SubElement(p_pr, f"{{{W}}}outlineLvl")
            node.set(f"{{{W}}}val", str(level - 1))
        else:
            document.add_heading(text, level=level)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_docx_with_table() -> bytes:
    document = Document()
    document.add_heading("Intro", level=1)
    document.add_paragraph("before table")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "cell A"
    table.cell(0, 1).text = "cell B"
    document.add_paragraph("after table")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_pptx(
    slides: Sequence[tuple[str | None, str | None]],
    sections: Sequence[tuple[str, Sequence[int]]] = (),
) -> bytes:
    """slides: (title, body); None leaves the placeholder empty. sections: (name, 1-based
    slide numbers)."""
    presentation = Presentation()
    for title, body in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        if title is not None:
            slide.shapes.title.text_frame.text = title
        if body is not None:
            slide.placeholders[1].text_frame.text = body
    if sections:
        ids = [s.get("id") for s in presentation.part._element.xpath("//p:sldIdLst/p:sldId")]
        ext_lst = etree.SubElement(presentation.part._element, f"{{{P}}}extLst")
        ext = etree.SubElement(ext_lst, f"{{{P}}}ext")
        ext.set("uri", "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
        section_list = etree.SubElement(ext, f"{{{P14}}}sectionLst", nsmap={"p14": P14})
        for index, (name, numbers) in enumerate(sections):
            section = etree.SubElement(section_list, f"{{{P14}}}section")
            section.set("name", name)
            section.set("id", f"{{00000000-0000-0000-0000-{index:012d}}}")
            id_list = etree.SubElement(section, f"{{{P14}}}sldIdLst")
            for number in numbers:
                etree.SubElement(id_list, f"{{{P14}}}sldId").set("id", ids[number - 1])
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


def with_extra_part(data: bytes, name: str, payload: bytes = b"x") -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, "w") as target:
        for info in source.infolist():
            target.writestr(info, source.read(info))
        target.writestr(name, payload)
    return out.getvalue()


def zip_bomb(data: bytes, inflated: int) -> bytes:
    """The same file plus one highly compressed part of `inflated` bytes."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source:
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                target.writestr(info, source.read(info))
            target.writestr("word/bomb.bin", b"\0" * inflated)
    return out.getvalue()
